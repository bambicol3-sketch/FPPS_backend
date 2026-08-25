import os
import uuid
from datetime import datetime

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.domains.pdf_to_ppt.application.port.pptx_builder_port import (
    BuiltPptx,
    PptxBuilderPort,
)
from app.domains.pdf_to_ppt.domain.value_object.slide_outline import SlideOutline

ACCENT = RGBColor(0x3A, 0x59, 0xD1)
DARK_TEXT = RGBColor(0x1A, 0x2B, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


class PptxBuilderImpl(PptxBuilderPort):
    """기본 테마(흰 배경 + 블루 강조색)로 title + bullet 슬라이드를 렌더링한다."""

    def build(
        self,
        deck_title: str,
        slides: list[SlideOutline],
        output_dir: str,
    ) -> BuiltPptx:
        prs = Presentation()
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)
        blank_layout = prs.slide_layouts[6]

        self._add_title_slide(prs, blank_layout, deck_title)
        for slide_outline in slides:
            self._add_content_slide(prs, blank_layout, slide_outline)

        os.makedirs(output_dir, exist_ok=True)
        file_name = (
            f"PDF2PPT_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:8]}.pptx"
        )
        file_path = os.path.join(output_dir, file_name)
        prs.save(file_path)

        return BuiltPptx(
            file_path=file_path,
            file_name=file_name,
            slide_count=len(prs.slides),
        )

    @staticmethod
    def _set_background(slide) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = WHITE

    def _add_title_slide(self, prs: Presentation, layout, title: str) -> None:
        slide = prs.slides.add_slide(layout)
        self._set_background(slide)

        box = slide.shapes.add_textbox(
            Inches(1), Inches(3), Inches(11.33), Inches(1.5)
        )
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title
        run.font.size = Pt(40)
        run.font.bold = True
        run.font.color.rgb = ACCENT

    def _add_content_slide(
        self, prs: Presentation, layout, slide_outline: SlideOutline
    ) -> None:
        slide = prs.slides.add_slide(layout)
        self._set_background(slide)

        title_box = slide.shapes.add_textbox(
            Inches(0.7), Inches(0.5), Inches(11.9), Inches(1)
        )
        title_tf = title_box.text_frame
        title_tf.word_wrap = True
        title_run = title_tf.paragraphs[0].add_run()
        title_run.text = slide_outline.title
        title_run.font.size = Pt(28)
        title_run.font.bold = True
        title_run.font.color.rgb = ACCENT

        underline = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.35), Inches(2.5), Pt(3)
        )
        underline.fill.solid()
        underline.fill.fore_color.rgb = ACCENT
        underline.line.fill.background()
        underline.shadow.inherit = False

        body_box = slide.shapes.add_textbox(
            Inches(0.9), Inches(1.7), Inches(11.5), Inches(5.3)
        )
        body_tf = body_box.text_frame
        body_tf.word_wrap = True
        bullets = slide_outline.bullets or ["(내용 없음)"]
        for i, bullet in enumerate(bullets):
            p = body_tf.paragraphs[0] if i == 0 else body_tf.add_paragraph()
            p.text = f"•  {bullet}"
            p.font.size = Pt(18)
            p.font.color.rgb = DARK_TEXT
            p.space_after = Pt(14)
