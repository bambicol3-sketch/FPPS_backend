import hashlib
import logging
import os
from pathlib import Path

from app.domains.template_rag.application.port.file_reader_port import (
    ExtractedFile,
    FileReaderPort,
)

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".csv",
    ".pdf",
    ".docx", ".doc",
    ".pptx", ".key",
    ".xlsx", ".xls",
    ".jpg", ".jpeg", ".png",
}


class FilesystemFileReader(FileReaderPort):
    def collect(self, folder_path: str) -> list[ExtractedFile]:
        if not folder_path or not os.path.isdir(folder_path):
            return []

        extracted: list[ExtractedFile] = []
        for path in sorted(Path(folder_path).rglob("*")):
            if not path.is_file():
                continue
            ext = path.suffix.lower()
            if ext not in SUPPORTED_EXTENSIONS:
                continue
            try:
                text = self._extract_text(path, ext)
                if not text.strip():
                    continue
                file_bytes = path.read_bytes()
                file_hash = hashlib.sha256(file_bytes).hexdigest()
                extracted.append(
                    ExtractedFile(
                        file_path=str(path.resolve()),
                        file_name=path.name,
                        text=text,
                        file_hash=file_hash,
                    )
                )
            except Exception as e:
                logger.warning("[FileReader] %s 추출 실패: %s", path, e)
        return extracted

    def _extract_text(self, path: Path, ext: str) -> str:
        if ext in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        if ext == ".csv":
            return self._extract_csv(path)
        if ext == ".pdf":
            return self._extract_pdf(path)
        if ext == ".docx":
            return self._extract_docx(path)
        if ext == ".doc":
            return self._extract_doc(path)
        if ext == ".pptx":
            return self._extract_pptx(path)
        if ext == ".key":
            return self._extract_key(path)
        if ext == ".xlsx":
            return self._extract_xlsx(path)
        if ext == ".xls":
            return self._extract_xls(path)
        if ext in {".jpg", ".jpeg", ".png"}:
            return self._extract_image(path, ext)
        return ""

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise RuntimeError(
                "pypdf 가 설치되어 있지 않습니다. pip install pypdf"
            ) from e
        reader = PdfReader(str(path))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)

    @staticmethod
    def _extract_docx(path: Path) -> str:
        try:
            import docx
        except ImportError as e:
            raise RuntimeError(
                "python-docx 가 설치되어 있지 않습니다. pip install python-docx"
            ) from e
        document = docx.Document(str(path))
        return "\n".join(p.text for p in document.paragraphs if p.text)

    @staticmethod
    def _extract_csv(path: Path) -> str:
        """CSV → 행별 탭 구분 문자열. BOM 제거 + 다양한 인코딩 시도."""
        import csv

        for enc in ("utf-8-sig", "utf-8", "cp949", "latin-1"):
            try:
                with open(path, "r", encoding=enc, newline="") as f:
                    parts: list[str] = []
                    reader = csv.reader(f)
                    for row in reader:
                        cells = [c for c in row if c is not None and c.strip()]
                        if cells:
                            parts.append("\t".join(cells))
                    return "\n".join(parts)
            except UnicodeDecodeError:
                continue
            except Exception as e:
                logger.warning("[FileReader] csv 추출 실패 (%s): %s", enc, e)
                return ""
        logger.warning("[FileReader] csv 인코딩 감지 실패: %s", path)
        return ""

    @staticmethod
    def _extract_image(path: Path, ext: str) -> str:
        """이미지(.jpg/.jpeg/.png) → OpenAI Vision 으로 텍스트·차트 설명 추출.

        - 이미지를 base64 인코딩해 Responses API 에 전달
        - 모델: settings.openai_finance_agent_model (기본 gpt-5-mini, 비전 지원)
        - 한국어 응답 (텍스트 + 핵심 설명)
        - API 키 없거나 호출 실패 시 빈 문자열 반환 (best-effort)
        """
        import base64

        try:
            from openai import OpenAI
            from app.infrastructure.config.settings import get_settings
        except Exception as e:
            logger.warning("[FileReader] OpenAI 의존성 실패 — 이미지 OCR skip: %s", e)
            return ""

        settings = get_settings()
        if not settings.openai_api_key:
            logger.warning(
                "[FileReader] OPENAI_API_KEY 없음 — 이미지 OCR skip: %s", path
            )
            return ""

        try:
            data = path.read_bytes()
        except Exception as e:
            logger.warning("[FileReader] 이미지 읽기 실패: %s", e)
            return ""

        mime_map = {
            ".jpg": "jpeg",
            ".jpeg": "jpeg",
            ".png": "png",
        }
        mime = mime_map.get(ext.lower(), "png")
        b64 = base64.b64encode(data).decode("ascii")
        data_url = f"data:image/{mime};base64,{b64}"

        prompt = (
            "이 이미지에 나타난 모든 텍스트(글자, 표 셀, 차트 라벨, 캡션, 수치) 를 그대로 추출하고, "
            "이미지가 차트/그래프/도식/사진이면 한 줄로 핵심 내용을 추가 설명해주세요. "
            "한국어 평문으로만 답변. 마크다운/JSON 없이."
        )
        try:
            client = OpenAI(api_key=settings.openai_api_key)
            response = client.responses.create(
                model=settings.openai_finance_agent_model,
                input=[
                    {
                        "role": "user",
                        "content": [
                            {"type": "input_text", "text": prompt},
                            {"type": "input_image", "image_url": data_url},
                        ],
                    }
                ],
            )
            text = (response.output_text or "").strip()
            logger.info(
                "[FileReader] 이미지 OCR 성공: %s (chars=%d)", path.name, len(text)
            )
            return text
        except Exception as e:
            logger.warning("[FileReader] 이미지 OCR 실패 (%s): %s", path, e)
            return ""

    @staticmethod
    def _extract_xlsx(path: Path) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise RuntimeError(
                "openpyxl 가 설치되어 있지 않습니다. pip install openpyxl"
            ) from e
        wb = load_workbook(str(path), data_only=True, read_only=True)
        parts: list[str] = []
        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                parts.append(f"[Sheet: {sheet_name}]")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    row_text = "\t".join(cells)
                    if row_text.strip():
                        parts.append(row_text)
                parts.append("")
        finally:
            wb.close()
        return "\n".join(parts)

    @staticmethod
    def _extract_xls(path: Path) -> str:
        try:
            import xlrd
        except ImportError:
            logger.warning(
                "[FileReader] .xls 추출 skip — xlrd 가 설치되어 있지 않음. "
                "pip install 'xlrd<2.0' (또는 파일을 .xlsx 로 변환): %s",
                path,
            )
            return ""
        try:
            book = xlrd.open_workbook(str(path))
        except Exception as e:
            logger.warning("[FileReader] .xls open 실패: %s", e)
            return ""
        parts: list[str] = []
        for sheet in book.sheets():
            parts.append(f"[Sheet: {sheet.name}]")
            for row_idx in range(sheet.nrows):
                row = sheet.row_values(row_idx)
                cells = [str(c) for c in row if c is not None and str(c).strip()]
                row_text = "\t".join(cells)
                if row_text.strip():
                    parts.append(row_text)
            parts.append("")
        return "\n".join(parts)

    @staticmethod
    def _extract_doc(path: Path) -> str:
        """구식 .doc 파일은 OLE 컨테이너. best-effort 텍스트 추출.

        정확도가 낮을 수 있어 가능하면 .docx 로 변환 후 사용 권장.
        """
        try:
            import olefile
        except ImportError:
            logger.warning(
                "[FileReader] .doc 추출 skip — olefile 미설치. "
                "pip install olefile (또는 .docx 변환 권장): %s",
                path,
            )
            return ""
        try:
            ole = olefile.OleFileIO(str(path))
        except Exception as e:
            logger.warning("[FileReader] .doc 컨테이너 open 실패: %s", e)
            return ""
        collected: list[str] = []
        try:
            for stream_path in ole.listdir():
                stream_name = "/".join(stream_path)
                # WordDocument 스트림에 본문 텍스트가 있으나 binary 인코딩 복잡.
                # 안전한 접근: 인쇄 가능 ASCII + 한글 UTF-16LE 디코딩 시도.
                try:
                    stream = ole.openstream(stream_path)
                    data = stream.read()
                except Exception:
                    continue
                # UTF-16LE 시도 (한글 포함)
                try:
                    decoded_u16 = data.decode("utf-16-le", errors="ignore")
                    cleaned = "".join(
                        ch for ch in decoded_u16
                        if ch.isprintable() or ch in "\n\t "
                    )
                    if len(cleaned.strip()) > 20:
                        collected.append(cleaned)
                        continue
                except Exception:
                    pass
                # ASCII 폴백
                try:
                    decoded = data.decode("latin-1", errors="ignore")
                    cleaned = "".join(
                        ch for ch in decoded
                        if ch.isprintable() or ch in "\n\t "
                    )
                    if len(cleaned.strip()) > 20:
                        collected.append(cleaned)
                except Exception:
                    pass
        finally:
            try:
                ole.close()
            except Exception:
                pass
        text = "\n".join(collected)
        # 중복 공백 정리
        import re as _re
        text = _re.sub(r"[ \t]{3,}", "  ", text)
        text = _re.sub(r"\n{3,}", "\n\n", text)
        return text

    @staticmethod
    def _extract_key(path: Path) -> str:
        """Apple Keynote (.key) 파일은 zip 아카이브이지만 본문이 .iwa(바이너리) 에 저장됨.

        best-effort: zip 내부 .xml/.plist/.txt/.html 텍스트만 추출.
        본문 정확도가 낮으면 Keynote.app 으로 .pptx 변환 후 ingest 권장.
        """
        import re as _re
        import zipfile

        try:
            zf = zipfile.ZipFile(str(path))
        except zipfile.BadZipFile:
            logger.warning("[FileReader] .key 가 zip 형식이 아님: %s", path)
            return ""
        parts: list[str] = []
        try:
            for name in zf.namelist():
                lower = name.lower()
                if not lower.endswith((".xml", ".plist", ".txt", ".html", ".rels")):
                    continue
                try:
                    raw = zf.read(name)
                except Exception:
                    continue
                try:
                    content = raw.decode("utf-8", errors="ignore")
                except Exception:
                    continue
                # XML/HTML 태그 제거
                stripped = _re.sub(r"<[^>]+>", " ", content)
                stripped = _re.sub(r"&[a-zA-Z]+;", " ", stripped)
                stripped = _re.sub(r"\s+", " ", stripped).strip()
                if len(stripped) > 10:
                    parts.append(stripped)
        finally:
            zf.close()
        return "\n\n".join(parts)

    @staticmethod
    def _extract_pptx(path: Path) -> str:
        try:
            from pptx import Presentation
        except ImportError as e:
            raise RuntimeError(
                "python-pptx 가 설치되어 있지 않습니다. pip install python-pptx"
            ) from e
        prs = Presentation(str(path))
        parts: list[str] = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        parts.append(line)
        return "\n".join(parts)
